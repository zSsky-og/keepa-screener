from collections import OrderedDict
import argparse
import csv
import hashlib
import json
import math
import os
import re
import sys
import time
from dataclasses import dataclass, replace
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile

import pandas as pd
import requests

MEMORY_LIMIT = 200
# KEEPA_API_BASE lets the test suite point the app at a fake Keepa server.
KEEPA_API_BASE = os.environ.get("KEEPA_API_BASE", "https://api.keepa.com").rstrip("/")
KEEPA_URL = KEEPA_API_BASE + "/product"
KEEPA_TOKEN_URL = KEEPA_API_BASE + "/token"
KEEPA_EPOCH_MIN = 21_564_000

DOMAINS = {"US": 1, "GB": 2, "DE": 3, "FR": 4, "JP": 5, "CA": 6,
           "IT": 8, "ES": 9, "IN": 10, "MX": 11}
AMAZON_TLD = {1: "com", 2: "co.uk", 3: "de", 4: "fr", 5: "co.jp", 6: "ca",
              8: "it", 9: "es", 10: "in", 11: "com.mx"}

AMAZON, NEW, SALES_RANK, COUNT_NEW, BUY_BOX = 0, 1, 3, 11, 18

CHECKS = ["amazon", "sales", "bsr", "offers", "roi", "suppression", "hazmat"]
FIELDS = ["asin", "code", "cost", "title", "pack"]

CACHE_DIR = Path(os.environ.get("KEEPA_CACHE_DIR") or Path(__file__).resolve().with_name(".keepa_cache"))
PROFILE_PATH = Path(os.environ.get("KEEPA_PROFILE_PATH") or Path(__file__).resolve().with_name("vendor_profiles.json"))

LIVE_KEYS = ["status", "source_row", "asin", "title", "supplier_title", "supplier_code",
             "unit_cost", "sell_price", "profit", "roi_pct", "monthly_sales", "bsr_now",
             "new_offers", "est_monthly_profit", "buybox_suppressed_pct_90d", "fail_reasons", "amazon_link",
             "brand", "pack_warning", "hazmat_info", "fees_estimated", "listing_pack"] + [f"pass_{c}" for c in CHECKS]

ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class Cancelled(Exception):
    pass


@dataclass
class Rules:
    min_monthly_sales: int = 50
    bsr_min: int = 1
    bsr_max: int = 50_000
    bsr_max_cv: float = 0.50
    min_offers: int = 3
    max_offers: int = 11
    amazon_max_instock_pct: float = 0.0
    min_roi: float = 20.0
    max_suppressed_pct: float = 30.0
    price_basis: str = "conservative"
    referral_fee_pct: float = 15.0
    fba_fee_default: float = 3.50
    prep_cost: float = 0.0
    inbound_ship: float = 0.0
    vat_rate: float = 0.20
    auto_pack: bool = False
    check_amazon: bool = True
    check_sales: bool = True
    check_bsr: bool = True
    check_offers: bool = True
    check_roi: bool = True
    check_suppression: bool = True
    check_hazmat: bool = True

    def enabled(self):
        return [c for c in CHECKS if getattr(self, f"check_{c}")]


QTY_ALIASES = ["qty", "quantity", "stock", "qty available", "available", "available qty",
               "units", "units available", "in stock", "stock qty", "stock level"]


@dataclass
class PreFilter:
    min_qty: float = None
    min_cost: float = None
    max_cost: float = None
    exclude_words: tuple = ()
    qty_column: str = None

    def active(self):
        return any(v is not None for v in (self.min_qty, self.min_cost, self.max_cost)) or bool(self.exclude_words)

    def describe(self):
        parts = []
        if self.min_qty is not None:
            parts.append(f"QTY >= {self.min_qty:g}")
        if self.min_cost is not None:
            parts.append(f"cost >= {self.min_cost:g}")
        if self.max_cost is not None:
            parts.append(f"cost <= {self.max_cost:g}")
        if self.exclude_words:
            parts.append("exclude: " + ", ".join(self.exclude_words))
        return "; ".join(parts) or "none"


ALIASES = {
    "asin": ["asin", "amazon asin"],
    "code": ["ean", "upc", "gtin", "barcode", "ean13", "ean 13", "upc/ean", "upc code",
             "ean code", "product code"],
    "cost": ["unit cost", "cost price", "wholesale price", "wholesale", "net cost", "your cost",
             "unit gbp", "price gbp", "unit price", "cost", "price", "gbp"],
    "title": ["product name", "product title", "title", "item description", "description", "name"],
    "pack": ["case size", "case pack", "case quantity", "pack qty", "pack size", "units per case",
             "qty per case", "case qty", "outer qty", "pack"],
}

HEADER_HINTS = ("ean", "upc", "asin", "gtin", "barcode", "descr", "title", "name", "sku", "cost",
                "price", "gbp", "case", "qty", "pack", "brand", "moq", "product")


def _norm(s):
    s = re.sub(r"[_\-.]+", " ", str(s).lower())
    s = re.sub(r"[^a-z0-9/ ]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def _s(v):
    if v is None or pd.isna(v):
        return None
    v = str(v).strip()
    return v or None


def excel_sheets(path):
    try:
        with pd.ExcelFile(path) as xl:
            return list(xl.sheet_names)
    except ImportError:
        raise ValueError("Old .xls files need the xlrd package (pip install xlrd). "
                         "Or open the file in Excel and save it as .xlsx") from None


def best_sheet(path):
    with pd.ExcelFile(path) as xl:
        best, best_key = xl.sheet_names[0], (-1, -1)
        for name in xl.sheet_names:
            try:
                raw = xl.parse(name, header=None, dtype=str)
            except Exception:
                continue
            if raw.empty:
                continue
            header = _header_row(raw)
            cells = [_norm(c) for c in raw.iloc[header].tolist() if _s(c)]
            score = sum(any(h in cell for h in HEADER_HINTS) for cell in cells)
            key = (score, len(raw) - header - 1)
            if key > best_key:
                best, best_key = name, key
        return best


def _detect_encoding(path):
    data = Path(path).read_bytes()
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            data.decode(encoding)
            return encoding
        except UnicodeDecodeError:
            continue
    return "latin-1"


def _detect_delimiter(path, encoding):
    with open(path, newline="", encoding=encoding) as handle:
        lines = [line for _, line in zip(range(200), handle)]
    best, best_width = ",", 0
    for sep in (",", ";", "\t", "|"):
        widths = sorted(len(r) for r in csv.reader(lines, delimiter=sep) if r)
        width = widths[len(widths) // 2] if widths else 0
        if width > best_width:
            best, best_width = sep, width
    return best


def _read_raw(path, sheet):
    ext = Path(path).suffix.lower()
    if ext in (".csv", ".txt"):
        encoding = _detect_encoding(path)
        sep = _detect_delimiter(path, encoding)
        with open(path, newline="", encoding=encoding) as handle:
            width = max((len(row) for row in csv.reader(handle, delimiter=sep)), default=1)
        raw = pd.read_csv(path, header=None, names=list(range(width)), dtype=str, sep=sep,
                          encoding=encoding, skip_blank_lines=False)
        return raw, encoding
    if ext in (".xlsx", ".xlsm", ".xls"):
        names = excel_sheets(path)
        if sheet in (None, ""):
            target = best_sheet(path)
        elif sheet in names:
            target = sheet
        elif str(sheet).isdigit() and int(sheet) < len(names):
            target = int(sheet)
        else:
            raise ValueError(f"Sheet {sheet!r} not found. Sheets in this file: {names}")
        raw = pd.read_excel(path, sheet_name=target, header=None, dtype=str)
        return raw, "excel"
    raise ValueError(f"Unsupported file type: {ext} (use .csv, .xlsx or .xls)")


def _header_row(raw):
    scan = min(len(raw), 30)
    best, best_score = None, 1
    for i in range(scan):
        cells = [_norm(c) for c in raw.iloc[i].tolist() if _s(c)]
        score = sum(any(h in cell for h in HEADER_HINTS) for cell in cells)
        if len(cells) >= 2 and score > best_score:
            best, best_score = i, score
    if best is not None:
        return best
    for i in range(scan):
        if sum(1 for c in raw.iloc[i].tolist() if _s(c)) >= 2:
            return i
    return 0


def load_catalog(path, sheet=None):
    raw, encoding = _read_raw(path, sheet)
    if raw.empty:
        raise ValueError("The file is empty")
    header = _header_row(raw)
    names, seen = [], {}
    for n, cell in enumerate(raw.iloc[header].tolist()):
        name = _s(cell) or f"column_{n + 1}"
        seen[name] = seen.get(name, 0) + 1
        names.append(name if seen[name] == 1 else f"{name}_{seen[name]}")
    df = raw.iloc[header + 1:].copy()
    df.columns = names
    df = df.apply(lambda col: col.map(lambda v: (v.strip() or None) if isinstance(v, str) else v))
    df = df.dropna(how="all")
    df = df.loc[:, df.notna().any(axis=0)]
    df.attrs.update(header_row=header + 1, encoding=encoding)
    return df


def _alias_in(alias, name):
    if len(alias) >= 4:
        return alias in name
    return re.search(r"\b" + re.escape(alias) + r"\b", name) is not None


def detect_columns(df, overrides=None):
    overrides = overrides or {}
    names = {c: _norm(c) for c in df.columns}
    mapping = {}
    used = {c for f, c in overrides.items() if f in ALIASES and c}
    for field, aliases in ALIASES.items():
        if field in overrides:
            chosen = overrides[field]
            if chosen and chosen not in df.columns:
                raise ValueError(f"Column {chosen!r} for '{field}' not found. "
                                 f"Columns are: {list(df.columns)}")
            mapping[field] = chosen or None
            continue
        free = {c: n for c, n in names.items() if c not in used}
        found = None
        for alias in aliases:
            found = next((c for c, n in free.items() if n == alias), None)
            if found:
                break
        if not found:
            for alias in aliases:
                if len(alias) < 3:
                    continue
                found = next((c for c, n in free.items() if _alias_in(alias, n)), None)
                if found:
                    break
        if found:
            used.add(found)
        mapping[field] = found
    return mapping


def guess_cost_basis(column):
    name = _norm(column or "")
    return "case" if re.search(r"\b(cases?|outers?|box(es)?|cartons?)\b", name) else "unit"


def resolve_basis(cost_basis, cost_column):
    return cost_basis if cost_basis in ("unit", "case") else guess_cost_basis(cost_column)


def profile_key(df):
    names = sorted(_norm(c) for c in df.columns)
    return hashlib.md5("|".join(names).encode()).hexdigest()[:12]


def _read_profiles():
    if not PROFILE_PATH.exists():
        return {}
    try:
        data = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    return data if isinstance(data, dict) else {}


def load_profile(df):
    saved = _read_profiles().get(profile_key(df))
    if not saved:
        return None
    stored = saved.get("mapping") or {}
    mapping = {f: (stored.get(f) or "") for f in FIELDS}
    mapping = {f: (c if c in df.columns else "") for f, c in mapping.items()}
    mapping["cost_basis"] = saved.get("cost_basis", "auto")
    return mapping


def save_profile(df, mapping, name, cost_basis):
    profiles = _read_profiles()
    profiles[profile_key(df)] = {"name": name, "cost_basis": cost_basis,
                                 "mapping": {f: mapping.get(f) or "" for f in FIELDS}}
    try:
        temp = PROFILE_PATH.with_suffix(".tmp")
        temp.write_text(json.dumps(profiles, indent=2), encoding="utf-8")
        temp.replace(PROFILE_PATH)
    except OSError:
        pass


def suggest_mapping(df):
    saved = load_profile(df)
    if saved:
        return saved, True
    return detect_columns(df), False


def valid_gtin(digits):
    if len(digits) not in (12, 13, 14):
        return False
    total = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(digits[:-1])))
    return (10 - total % 10) % 10 == int(digits[-1])


def clean_asin(v):
    v = (_s(v) or "").upper()
    return v if re.fullmatch(r"[A-Z0-9]{10}", v) else None


def clean_code(v):
    v = _s(v)
    if not v:
        return None
    if re.fullmatch(r"\d+(\.\d+)?E\+?\d+", v, re.I):
        try:
            v = str(int(Decimal(v)))
        except (InvalidOperation, ValueError):
            return None
    digits = re.sub(r"\D", "", v.split(".")[0])
    if len(digits) < 8 or len(digits) > 14:
        return None
    if len(digits) == 8:
        return digits if valid_gtin(digits.zfill(12)) else None
    if len(digits) < 12:
        digits = digits.zfill(12)
    return digits if valid_gtin(digits) else None


def clean_num(v):
    v = _s(v)
    if not v:
        return None
    v = re.sub(r"[^\d.,\-]", "", v)
    if "," in v and "." in v:
        if v.rfind(",") > v.rfind("."):
            v = v.replace(".", "").replace(",", ".")
        else:
            v = v.replace(",", "")
    elif "," in v:
        if re.fullmatch(r"-?\d{1,3}(,\d{3})+", v):
            v = v.replace(",", "")
        else:
            v = v.replace(",", ".")
    try:
        return float(v)
    except ValueError:
        return None


def build_items(df, cols, cost_basis=None):
    basis = resolve_basis(cost_basis, cols.get("cost"))

    def get(row, field):
        return row.get(cols[field]) if cols.get(field) else None

    items = []
    for index, row in df.iterrows():
        raw_code = _s(get(row, "code"))
        code = clean_code(raw_code)
        pack = clean_num(get(row, "pack"))
        pack = pack if pack and pack > 0 else 1
        cost = clean_num(get(row, "cost"))
        unit_cost = cost / pack if basis == "case" and cost is not None else cost
        items.append({
            "source_row": int(index) + 1,
            "asin": clean_asin(get(row, "asin")),
            "code": code,
            "bad_code": bool(raw_code) and code is None,
            "title": _s(get(row, "title")),
            "pack": pack,
            "case_size": pack if cols.get("pack") else None,
            "cost": cost,
            "unit_cost": unit_cost,
            "src": {f"src_{k}": _s(v) for k, v in row.to_dict().items()},
        })
    return items


def find_qty_column(df, cols, override=None):
    if override:
        if override not in df.columns:
            raise ValueError(f"QTY column {override!r} not found. Columns are: {list(df.columns)}")
        return override
    used = {c for c in cols.values() if c}
    names = {c: _norm(c) for c in df.columns if c not in used}
    for alias in QTY_ALIASES:
        found = next((c for c, n in names.items() if n == alias), None)
        if found:
            return found
    return None


def _word_pattern(word):
    word = word.strip().lower()
    if not word:
        return None
    if word.endswith("*"):
        return re.compile(r"\b" + re.escape(word[:-1]), re.I)
    return re.compile(r"\b" + re.escape(word) + r"\b", re.I)


def prefilter_reason(item, prefilter, qty_column=None, patterns=None):
    if not prefilter or not prefilter.active():
        return None
    cost = item.get("unit_cost")
    if prefilter.min_cost is not None and cost is not None and cost < prefilter.min_cost:
        return f"Cost {cost:g} below minimum {prefilter.min_cost:g}"
    if prefilter.max_cost is not None and cost is not None and cost > prefilter.max_cost:
        return f"Cost {cost:g} above maximum {prefilter.max_cost:g}"
    if prefilter.min_qty is not None and qty_column:
        qty = clean_num(item["src"].get(f"src_{qty_column}"))
        if qty is not None and qty < prefilter.min_qty:
            return f"QTY {qty:g} below minimum {prefilter.min_qty:g}"
    title = item.get("title") or ""
    for word, pattern in (patterns or []):
        if pattern.search(title):
            return f"Name contains excluded word '{word}'"
    return None


def count_lookups(items):
    asins = {i["asin"] for i in items if i["asin"]}
    codes = {i["code"] for i in items if not i["asin"] and i["code"]}
    skipped = sum(1 for i in items if not i["asin"] and not i["code"])
    return len(asins), len(codes), skipped


def token_status(api_key):
    try:
        r = requests.get(KEEPA_TOKEN_URL, params={"key": api_key}, timeout=20)
    except requests.RequestException as exc:
        raise RuntimeError(f"Could not reach Keepa ({exc.__class__.__name__})") from None
    try:
        data = r.json() if r.content else {}
    except ValueError:
        data = {}
    if r.status_code != 200:
        raise RuntimeError(str(data.get("error") or f"Keepa returned HTTP {r.status_code}").replace(api_key, "***"))
    return {"left": data.get("tokensLeft"), "rate": data.get("refillRate")}


def _norm_code(c):
    return re.sub(r"\D", "", str(c)).lstrip("0")


def _product_codes(p):
    out = set()
    for key in ("eanList", "upcList", "gtinList"):
        for c in p.get(key) or []:
            out.add(_norm_code(c))
    return out


class TokenBucket:
    """What we last heard from Keepa about the account's tokens.

    Shared by every client in one run, so the first and second pass never
    fight over the same tokens or both hammer Keepa with 429s.
    """

    def __init__(self):
        self.left = None
        self.rate = None
        self.refill_at = None

    def update(self, data):
        if not isinstance(data, dict):
            return
        if data.get("tokensLeft") is not None:
            self.left = data["tokensLeft"]
        if data.get("refillRate"):
            self.rate = data["refillRate"]
        if data.get("refillIn") is not None:
            self.refill_at = time.time() + max(0, data["refillIn"]) / 1000.0

    def seconds_until_positive(self):
        """Rough time until the balance is above zero again."""
        if self.left is None or self.left > 0:
            return 0.0
        first = max(0.0, (self.refill_at or time.time()) - time.time())
        deficit = 1 - self.left
        if self.rate and deficit > self.rate:
            first += math.ceil((deficit - self.rate) / self.rate) * 60
        return max(1.0, first + 0.5)


class KeepaClient:
    """Looks products up on Keepa strictly one at a time.

    Each lookup is one request. The result is handed back (on_batch) the moment
    it arrives, so the caller can score it and show the row straight away.
    Keepa allows a request whenever the balance is above zero (it may go
    negative), so we only ever wait when the balance is zero or below.
    """

    def __init__(self, key, domain, buybox=True, batch=1, cache_dir=CACHE_DIR,
                 cache_hours=24, log=print, stop=None, progress=None, offline=False, on_tokens=None,
                 pause=None, on_wait=None, bucket=None):
        self.key, self.domain, self.buybox = key, domain, buybox
        self.batch = 1  # kept for backwards compatibility, lookups are always one by one
        self.on_tokens = on_tokens
        self.on_wait = on_wait or (lambda reason, until=None: None)
        self.offline = offline
        self.cache_dir, self.cache_hours, self.log = Path(cache_dir), cache_hours, log
        self.stop = stop
        self.pause = pause
        self.progress = progress or (lambda *args: None)
        self.cache_dir.mkdir(parents=True, exist_ok=True)
        self.session = requests.Session()
        self.bucket = bucket or TokenBucket()
        self.per_product = 3 if buybox else 1
        self.memory = OrderedDict()

    @property
    def tokens_left(self):
        return self.bucket.left

    def _check_stop(self):
        if self.stop and self.stop():
            raise Cancelled()
        if self.pause and self.pause():
            self.on_wait("paused")
            while self.pause():
                if self.stop and self.stop():
                    raise Cancelled()
                time.sleep(0.2)
            self.on_wait(None)

    def _sleep(self, seconds):
        end = time.time() + seconds
        while time.time() < end:
            self._check_stop()
            time.sleep(min(0.25, max(0.0, end - time.time())))

    def _cache_path(self, kind, ident):
        raw = f"{self.domain}|{kind}|{ident}|{int(self.buybox)}"
        return self.cache_dir / f"{hashlib.md5(raw.encode()).hexdigest()}.json"

    def _cache_get(self, kind, ident):
        if (kind, ident) in self.memory:
            return self.memory[(kind, ident)]
        if self.cache_hours <= 0:
            return None
        path = self._cache_path(kind, ident)
        try:
            if time.time() - path.stat().st_mtime >= self.cache_hours * 3600:
                return None
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            return None
        return data if isinstance(data, list) else None

    def is_cached(self, kind, ident):
        return self._cache_get(kind, ident) is not None

    def _cache_put(self, kind, ident, products):
        self.memory[(kind, ident)] = products
        self.memory.move_to_end((kind, ident))
        while len(self.memory) > MEMORY_LIMIT:
            self.memory.popitem(last=False)
        path = self._cache_path(kind, ident)
        temp = path.with_suffix(".tmp")
        try:
            with open(temp, "w", encoding="utf-8") as handle:
                handle.write(json.dumps(products))
                handle.flush()
                os.fsync(handle.fileno())
            temp.replace(path)
        except OSError as exc:
            self.log(f"Could not write the cache file ({exc.__class__.__name__}), continuing without it")

    def _tokens_changed(self):
        if self.on_tokens and self.bucket.left is not None:
            self.on_tokens(self.bucket.left)

    def _refresh_tokens(self):
        """Ask Keepa for the balance. The /token call itself costs nothing."""
        try:
            r = self.session.get(KEEPA_TOKEN_URL, params={"key": self.key}, timeout=20)
            if r.status_code == 200:
                self.bucket.update(r.json())
                self._tokens_changed()
                return True
        except (requests.RequestException, ValueError):
            pass
        return False

    def _wait_for_tokens(self):
        announced = False
        while self.bucket.left is not None and self.bucket.left <= 0:
            seconds = self.bucket.seconds_until_positive()
            if not announced:
                self.log(f"Out of Keepa tokens ({self.bucket.left} left), waiting about {seconds:.0f}s for the refill")
                announced = True
            self.on_wait("tokens", time.time() + seconds)
            self._sleep(seconds)
            if not self._refresh_tokens():
                break  # cannot check, just try the request and let Keepa decide
        if announced:
            self.on_wait(None)
            self.log(f"Tokens refilled ({self.bucket.left} left), carrying on")

    def _request(self, kind, ident):
        params = {"key": self.key, "domain": self.domain, "stats": 90,
                  "days": 90, "history": 1, kind: ident}
        if self.buybox:
            params["buybox"] = 1
        failures = 0
        limit_waits = 0
        while True:
            self._check_stop()
            self._wait_for_tokens()
            self._check_stop()
            try:
                r = self.session.get(KEEPA_URL, params=params, timeout=120)
            except requests.RequestException as exc:
                failures += 1
                if failures > 5:
                    raise RuntimeError(f"Could not reach Keepa ({exc.__class__.__name__})") from None
                pause = min(60, 5 * 2 ** (failures - 1))
                self.log(f"Network problem ({exc.__class__.__name__}), retrying in {pause}s")
                self.on_wait("network", time.time() + pause)
                self._sleep(pause)
                self.on_wait(None)
                continue
            try:
                data = r.json()
            except ValueError:
                data = {}
            if not isinstance(data, dict):
                data = {}
            self.bucket.update(data)
            self._tokens_changed()
            if r.status_code == 429:
                limit_waits += 1
                if limit_waits > 1000:
                    raise RuntimeError("Keepa kept refusing requests (HTTP 429) for too long")
                if self.bucket.left is None or self.bucket.left > 0:
                    # Keepa said no although the balance looks positive: short pause, try again.
                    pause = max(1.0, min(10.0, (data.get("refillIn") or 5000) / 1000.0))
                    self.on_wait("tokens", time.time() + pause)
                    self._sleep(pause)
                    self.on_wait(None)
                continue  # with a balance of zero or less, _wait_for_tokens does the waiting
            if r.status_code >= 500 or (r.status_code == 200 and "products" not in data):
                failures += 1
                if failures > 5:
                    raise RuntimeError(f"Keepa server error {r.status_code}" if r.status_code >= 500 else "Keepa sent a reply that is not product data")
                pause = min(60, 5 * 2 ** (failures - 1))
                self.log(f"Keepa server error {r.status_code}, retrying in {pause}s" if r.status_code >= 500 else f"Keepa sent a reply that is not product data, retrying in {pause}s")
                self.on_wait("network", time.time() + pause)
                self._sleep(pause)
                self.on_wait(None)
                continue
            if r.status_code != 200:
                detail = str(data.get("error") or r.text[:300])
                if self.key:
                    detail = detail.replace(self.key, "***")
                if r.status_code in (400, 401, 402, 403) and "key" in detail.lower():
                    raise RuntimeError(f"Keepa rejected the API key: {detail}")
                raise RuntimeError(f"Keepa API error {r.status_code}: {detail}")
            return data

    @staticmethod
    def _match(kind, ident, products):
        if kind == "asin":
            return [p for p in products if str(p.get("asin", "")).upper() == ident]
        return [p for p in products if _norm_code(ident) in _product_codes(p)]

    def lookup(self, kind, ident):
        """One product, one request. Returns the list of matching Keepa products."""
        data = self._request(kind, ident)
        products = [p for p in (data.get("products") or []) if isinstance(p, dict)]
        found = self._match(kind, ident, products)
        if not found and kind == "code":
            found = products  # Keepa matched the code even though the lists do not show it
        self._cache_put(kind, ident, found)
        name = "ASIN" if kind == "asin" else "barcode"
        self.log(f"Fetched {name} {ident}: {len(found)} product(s) (tokens left: {self.bucket.left})")
        return found

    def fetch(self, kind, ids, label="", on_batch=None):
        """Looks every id up one by one, calling on_batch({id: products}) after each."""
        unique = list(dict.fromkeys(ids))
        out = {}
        total = len(unique)
        cached = missing = 0
        self.progress(label, 0, total)
        for done, ident in enumerate(unique, 1):
            self._check_stop()
            found = self._cache_get(kind, ident)
            if found is not None:
                cached += 1
            elif self.offline:
                missing += 1
                self.progress(label, done, total)
                continue
            else:
                found = self.lookup(kind, ident)
            out[ident] = found
            if on_batch:
                on_batch({ident: found})
            self.progress(label, done, total)
        if cached:
            self.log(f"{cached} {kind} lookups served from cache")
        if missing:
            self.log(f"{missing} {kind} lookups are not in the cache and were not fetched (cache-only mode)")
        return out


def val(arr, i):
    try:
        v = arr[i]
    except (TypeError, IndexError):
        return None
    return None if v is None or v < 0 else v


def series_stats(arr, days=90):
    if not arr or len(arr) < 2:
        return None
    now = int(time.time() / 60) - KEEPA_EPOCH_MIN
    cutoff = now - days * 1440
    points = [(arr[i], arr[i + 1]) for i in range(0, len(arr) - 1, 2)]
    segments = []
    for j, (t, v) in enumerate(points):
        t_end = points[j + 1][0] if j + 1 < len(points) else now
        t_start = max(t, cutoff)
        if t_end <= cutoff or t_end <= t_start or v is None or v < 0:
            continue
        segments.append((v, t_end - t_start))
    total = sum(w for _, w in segments)
    if total <= 0:
        return None
    mean = sum(v * w for v, w in segments) / total
    std = math.sqrt(sum(w * (v - mean) ** 2 for v, w in segments) / total)
    return {"mean": mean, "std": std, "cv": (std / mean) if mean else None}


def hazmat_info(p):
    found = [m for m in (p.get("hazardousMaterials") or []) if m]
    if not found:
        return None
    labels = {"proper_shipping_name": "", "united_nations_regulatory_id": "",
              "transportation_regulatory_class": "class ", "regulatory_packing_group": "group "}
    parts = []
    for m in found:
        if not isinstance(m, dict):
            parts.append(str(m))
            continue
        aspect, value = str(m.get("aspect") or ""), str(m.get("value") or "")
        if aspect in labels and value:
            parts.append(labels[aspect] + value)
        elif aspect == "exception":
            if value.startswith("GB_"):
                parts.append("GB " + value[3:])
        elif aspect or value:
            parts.append(f"{aspect} {value}".strip())
    return "; ".join(dict.fromkeys(parts)) or "flagged"


def suppression_stats(p, days=90):
    history = p.get("buyBoxSellerIdHistory")
    if isinstance(history, list) and len(history) >= 2:
        now = int(time.time() / 60) - KEEPA_EPOCH_MIN
        cutoff = now - days * 1440
        points = []
        for i in range(0, len(history) - 1, 2):
            try:
                points.append((int(history[i]), str(history[i + 1])))
            except (TypeError, ValueError):
                continue
        total = suppressed = 0
        for j, (t, seller) in enumerate(points):
            t_end = points[j + 1][0] if j + 1 < len(points) else now
            t_start = max(t, cutoff)
            if t_end <= t_start:
                continue
            total += t_end - t_start
            if seller == "-1":
                suppressed += t_end - t_start
        if total > 0:
            return suppressed / total * 100.0, "Buy Box seller history"
    oos = (p.get("stats") or {}).get("outOfStockPercentage90")
    missing = val(oos, BUY_BOX)
    if missing is not None:
        return float(missing), "Buy Box price missing % (approximate)"
    return None, None


PACK_PATTERNS = [re.compile(x, re.I) for x in (r"\bpack\s+of\s+(\d+)", r"^\s*(\d+)\s*x\s+\S", r"\bx\s*\(?(\d+)\)?\s*$")]


def listing_pack(p):
    qty = p.get("packageQuantity")
    if isinstance(qty, (int, float)) and not isinstance(qty, bool) and qty > 1:
        return int(qty)
    return 1


def pack_warning(p, adjusted=1):
    notes = []
    qty = listing_pack(p)
    if qty > 1:
        notes.append(f"cost x{qty} (Keepa pack qty)" if adjusted > 1 else f"Keepa pack qty {qty}")
    title = str(p.get("title") or "")
    for pattern in PACK_PATTERNS:
        m = pattern.search(title)
        if m and int(m.group(1)) > 1:
            if int(m.group(1)) != adjusted:
                notes.append(f'title says "{m.group(0).strip()}"')
            break
    return "; ".join(notes) or None


def optimistic_roi(p, unit_cost, rules, domain):
    if unit_cost is None or unit_cost <= 0:
        return None
    if rules.auto_pack:
        unit_cost = unit_cost * listing_pack(p)
    stats = p.get("stats") or {}
    prices = [val(stats.get(key), index) for key in ("current", "avg90") for index in (NEW, BUY_BOX)]
    box = stats.get("buyBoxPrice")
    prices.append(None if box is None or box < 0 else box)
    prices = [x for x in prices if x]
    if not prices:
        return None
    divisor = 1 if domain == DOMAINS["JP"] else 100
    sell = max(prices) * 1.15 / divisor
    fee_info = p.get("fbaFees") or {}
    referral = p.get("referralFeePercentage") or p.get("referralFeePercent") or rules.referral_fee_pct
    fba = fee_info.get("pickAndPackFee")
    fba = fba / divisor if fba else None
    if fba is None or fba >= sell:
        fba = rules.fba_fee_default
    return calc_roi(sell, unit_cost, referral, fba, rules)[1]


def calc_roi(sell, unit_cost, referral_pct, fba_fee, rules):
    net_sell = sell / (1 + rules.vat_rate) if rules.vat_rate else sell
    fixed = sell * referral_pct / 100.0 + fba_fee + rules.prep_cost + rules.inbound_ship
    profit = net_sell - fixed - unit_cost
    roi = (profit / unit_cost * 100.0) if unit_cost and unit_cost > 0 else None
    max_cost = (net_sell - fixed) / (1 + rules.min_roi / 100.0)
    return profit, roi, max_cost


def _round(x, n=2):
    return None if x is None else round(x, n)


def _fmt(x):
    return "n/a" if x is None else f"{x:,.0f}"


def evaluate(p, unit_cost, rules, domain):
    stats = p.get("stats") or {}
    cur, avg30, avg90 = stats.get("current"), stats.get("avg30"), stats.get("avg90")
    divisor = 1 if domain == DOMAINS["JP"] else 100

    def money(cents):
        return None if cents is None else cents / divisor

    lp = listing_pack(p) if rules.auto_pack else 1
    cost = unit_cost * lp if unit_cost is not None and unit_cost > 0 else unit_cost
    base = {"asin": p.get("asin"), "title": p.get("title"), "brand": p.get("brand"),
            "pkg_qty": p.get("packageQuantity"), "pack_warning": pack_warning(p, lp),
            "keepa_link": f"https://keepa.com/#!product/{domain}-{p.get('asin')}",
            "amazon_link": f"https://www.amazon.{AMAZON_TLD.get(domain, 'com')}/dp/{p.get('asin')}"}
    if not stats or not cur:
        return {**base, "status": "NO_DATA", "fail_reasons": "No Keepa stats returned for this product"}

    amazon_now = val(cur, AMAZON)
    availability = p.get("availabilityAmazon")
    oos = stats.get("outOfStockPercentage90")
    amazon_oos = oos[AMAZON] if isinstance(oos, list) and len(oos) > AMAZON else None
    amazon_pct = (100 - amazon_oos) if amazon_oos is not None and amazon_oos >= 0 else None
    notes = []
    if amazon_now is not None:
        notes.append("Amazon has a live price")
    if availability in (0, 1, 3, 4):
        notes.append("Amazon offer flagged available")
    if stats.get("buyBoxIsAmazon"):
        notes.append("Amazon owns the Buy Box")
    if amazon_pct is not None and amazon_pct > rules.amazon_max_instock_pct:
        notes.append(f"Amazon in stock {amazon_pct:.0f}% of last 90d")
    pass_amazon = not notes
    reasons_amazon = ["Amazon present (" + ", ".join(notes) + ")"] if notes else []

    monthly = p.get("monthlySold")
    sales_source = "Keepa monthlySold"
    if monthly is None or monthly < 0:
        monthly = stats.get("salesRankDrops30")
        if monthly is not None and monthly < 0:
            monthly = None  # Keepa uses -1 for "no data"
        sales_source = "rank drops 30d (proxy)" if monthly is not None else None
    pass_sales = monthly is not None and monthly >= rules.min_monthly_sales
    reasons_sales = []
    if monthly is None:
        reasons_sales.append("No monthly sales data")
    elif not pass_sales:
        reasons_sales.append(f"Monthly sales {monthly} < {rules.min_monthly_sales}")

    bsr_now, bsr30, bsr90 = val(cur, SALES_RANK), val(avg30, SALES_RANK), val(avg90, SALES_RANK)
    csv_data = p.get("csv") or []
    series = series_stats(csv_data[SALES_RANK] if len(csv_data) > SALES_RANK else None)
    cv = series["cv"] if series else None
    known = [x for x in (bsr_now, bsr30, bsr90) if x is not None]
    in_range = bool(known) and all(rules.bsr_min <= x <= rules.bsr_max for x in known)
    steady = cv is not None and cv <= rules.bsr_max_cv
    pass_bsr = in_range and steady
    reasons_bsr = []
    if not known:
        reasons_bsr.append("No BSR data")
    elif not in_range:
        reasons_bsr.append(f"BSR outside {rules.bsr_min:,}-{rules.bsr_max:,} "
                           f"(now {_fmt(bsr_now)}, 30d avg {_fmt(bsr30)}, 90d avg {_fmt(bsr90)})")
    elif not steady:
        cv_text = "n/a" if cv is None else f"{cv:.2f}"
        reasons_bsr.append(f"BSR not steady (CV {cv_text} > {rules.bsr_max_cv})")

    offers = val(cur, COUNT_NEW)
    if offers is None:
        offers = stats.get("totalOfferCount")
        if offers is not None and offers < 0:
            offers = None
    pass_offers = offers is not None and rules.min_offers <= offers <= rules.max_offers
    reasons_offers = []
    if offers is None:
        reasons_offers.append("No offer count")
    elif not pass_offers:
        reasons_offers.append(f"New offers {offers} outside {rules.min_offers}-{rules.max_offers}")

    buybox_now = stats.get("buyBoxPrice")
    buybox_now = None if buybox_now is None or buybox_now < 0 else buybox_now
    buybox_90, new_now, new_90 = val(avg90, BUY_BOX), val(cur, NEW), val(avg90, NEW)

    def usable(*values):
        return [x for x in values if x]

    if rules.price_basis == "current":
        candidates = usable(buybox_now) or usable(new_now)
    elif rules.price_basis == "avg90":
        candidates = usable(buybox_90) or usable(new_90)
    else:
        candidates = usable(buybox_now, buybox_90) or usable(new_now, new_90)
    sell = money(min(candidates)) if candidates else None

    fee_info = p.get("fbaFees") or {}
    referral_pct = p.get("referralFeePercentage") or p.get("referralFeePercent")
    estimated = []
    if not referral_pct:
        referral_pct = rules.referral_fee_pct
        estimated.append("referral")
    fba = money(fee_info["pickAndPackFee"]) if fee_info.get("pickAndPackFee") else None
    if fba is None or (sell is not None and fba >= sell):
        fba = rules.fba_fee_default
        estimated.append("FBA")

    profit = roi = max_cost = None
    reasons_roi = []
    if sell is None:
        reasons_roi.append("No sell price available")
    elif unit_cost is None or unit_cost <= 0:
        reasons_roi.append("No supplier cost in file")
    else:
        profit, roi, max_cost = calc_roi(sell, cost, referral_pct, fba, rules)
    pass_roi = roi is not None and roi >= rules.min_roi
    if roi is not None and not pass_roi:
        reasons_roi.append(f"ROI {roi:.1f}% < {rules.min_roi:g}%")

    hazmat = hazmat_info(p)
    pass_hazmat = hazmat is None
    reasons_hazmat = [f"HAZMAT product ({hazmat})"] if hazmat else []

    suppressed, suppression_source = suppression_stats(p)
    pass_suppression = suppressed is None or suppressed <= rules.max_suppressed_pct
    reasons_suppression = []
    if suppressed is not None and not pass_suppression:
        reasons_suppression.append(
            f"Buy Box suppressed {suppressed:.1f}% of last 90d > {rules.max_suppressed_pct:g}%")

    units = monthly / (offers + 1) if monthly is not None and offers is not None else None
    est_profit = units * profit if units is not None and profit is not None else None

    rows = [("amazon", pass_amazon, reasons_amazon), ("sales", pass_sales, reasons_sales),
            ("bsr", pass_bsr, reasons_bsr), ("offers", pass_offers, reasons_offers),
            ("roi", pass_roi, reasons_roi), ("suppression", pass_suppression, reasons_suppression),
            ("hazmat", pass_hazmat, reasons_hazmat)]
    active = [(name, ok, why) for name, ok, why in rows if getattr(rules, f"check_{name}")]
    passed = all(ok for _, ok, _ in active)
    reasons = [text for _, _, why in active for text in why]
    flag = {name: (ok if getattr(rules, f"check_{name}") else None) for name, ok, _ in rows}

    return {
        **base,
        "status": "PASS" if passed else "FAIL",
        "fail_reasons": "; ".join(reasons),
        "unit_cost": _round(unit_cost), "listing_pack": lp, "cost_per_listing": _round(cost),
        "sell_price": _round(sell), "profit": _round(profit),
        "roi_pct": _round(roi, 1), "max_cost_for_min_roi": _round(max_cost / lp if max_cost is not None else None),
        "est_units_per_seller": _round(units, 1), "est_monthly_profit": _round(est_profit, 0),
        "referral_fee_pct": referral_pct, "fba_fee": _round(fba),
        "fees_estimated": ", ".join(estimated) or None,
        "monthly_sales": monthly, "sales_source": sales_source,
        "bsr_now": bsr_now, "bsr_avg30": _round(bsr30, 0), "bsr_avg90": _round(bsr90, 0),
        "bsr_cv": _round(cv), "new_offers": offers,
        "buybox_suppressed_pct_90d": _round(suppressed, 1), "suppression_source": suppression_source,
        "hazmat_info": hazmat,
        "amazon_instock_pct_90d": _round(amazon_pct, 0),
        "pass_amazon": flag["amazon"], "pass_sales": flag["sales"], "pass_bsr": flag["bsr"],
        "pass_offers": flag["offers"], "pass_roi": flag["roi"],
        "pass_suppression": flag["suppression"], "pass_hazmat": flag["hazmat"],
    }


FRONT = ["asin", "status", "fail_reasons", "pack_warning", "fees_estimated", "title", "brand", "supplier_code",
         "unit_cost", "listing_pack", "cost_per_listing", "sell_price", "profit", "roi_pct", "est_monthly_profit", "buybox_suppressed_pct_90d",
         "hazmat_info", "case_size", "profit_per_case",
         "max_cost_for_min_roi", "monthly_sales", "bsr_now", "bsr_avg30", "bsr_avg90", "bsr_cv",
         "new_offers", "pkg_qty", "pack_qty", "keepa_link", "amazon_link"]


def _style(ws):
    from openpyxl.styles import Alignment, Font, PatternFill
    header_font = Font(name="Arial", bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10)
    fill = PatternFill("solid", fgColor="1F3864")
    for cell in ws[1]:
        cell.font, cell.fill = header_font, fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
    for column in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in list(column)[:200])
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 55)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _link_asins(ws):
    from openpyxl.styles import Font
    headers = {cell.value: cell.column for cell in ws[1]}
    asin_col, link_col = headers.get("asin"), headers.get("amazon_link")
    if not asin_col or not link_col:
        return
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    for row in range(2, ws.max_row + 1):
        link = ws.cell(row=row, column=link_col).value
        if link:
            cell = ws.cell(row=row, column=asin_col)
            cell.hyperlink = str(link)
            cell.font = link_font


def _clean_cell(v):
    if not isinstance(v, str):
        return v
    v = ILLEGAL_CHARS.sub("", v)
    return " " + v if v.startswith("=") else v


def write_report(rows, rules, domain_name, path, prefilter=None):
    df = pd.DataFrame(rows)
    for col in ("status", "est_monthly_profit", "roi_pct"):
        if col not in df.columns:
            df[col] = None
    for col in ("est_monthly_profit", "roi_pct"):
        df[col] = pd.to_numeric(df[col], errors="coerce")
    ordered = ([c for c in FRONT if c in df.columns]
               + [c for c in df.columns if c not in FRONT and not c.startswith("src_")]
               + [c for c in df.columns if c.startswith("src_")])
    df = df[ordered]
    df.columns = [_clean_cell(c) for c in df.columns]
    df = df.apply(lambda col: col.map(_clean_cell))
    passed = df[df.status == "PASS"].sort_values(["est_monthly_profit", "roi_pct"],
                                                 ascending=False, na_position="last")
    failed = df[df.status != "PASS"]
    matched = df[df.status.isin(["PASS", "FAIL"])]

    def count_failed(col):
        return int((matched[col] == False).sum()) if col in matched else 0

    summary = [
        ("Marketplace", domain_name),
        ("Rows in output", len(df)),
        ("PASSED all rules", len(passed)),
        ("Passed with a pack warning", int(passed["pack_warning"].notna().sum()) if "pack_warning" in passed else 0),
        ("Failed", int((df.status == "FAIL").sum())),
        ("Not found on Keepa / no data", int(df.status.isin(["NOT_FOUND", "NO_DATA"]).sum())),
        ("No usable ASIN/UPC/EAN in row", int((df.status == "NO_ID").sum())),
        ("Skipped before Keepa (input filters)", int((df.status == "SKIPPED").sum())),
        ("Not in cache (cache-only mode)", int((df.status == "NOT_CACHED").sum())),
        ("Stopped before the final check", int((df.status == "STOPPED").sum())),
        ("Not checked (run stopped first)", int((df.status == "NOT_CHECKED").sum())),
        ("", ""),
        ("Failed: Amazon present", count_failed("pass_amazon")),
        ("Failed: monthly sales", count_failed("pass_sales")),
        ("Failed: BSR range/steadiness", count_failed("pass_bsr")),
        ("Failed: offer count", count_failed("pass_offers")),
        ("Failed: ROI", count_failed("pass_roi")),
        ("Failed: Buy Box suppression", count_failed("pass_suppression")),
        ("Failed: HAZMAT", count_failed("pass_hazmat")),
        ("", ""),
        ("RULES USED", ""),
        ("Checks enabled", ", ".join(rules.enabled())),
        ("Min monthly sales", rules.min_monthly_sales),
        ("BSR range", f"{rules.bsr_min:,} - {rules.bsr_max:,}"),
        ("BSR steadiness (max CV, 90d)", rules.bsr_max_cv),
        ("New offers range", f"{rules.min_offers} - {rules.max_offers}"),
        ("Amazon max in-stock % (90d)", rules.amazon_max_instock_pct),
        ("Min ROI %", rules.min_roi),
        ("Max Buy Box suppression % (90d)", rules.max_suppressed_pct),
        ("Sell price basis", rules.price_basis),
        ("Prep cost / unit", rules.prep_cost),
        ("Inbound ship / unit", rules.inbound_ship),
        ("VAT rate", rules.vat_rate),
        ("Cost multiplied by Keepa pack qty", "yes" if rules.auto_pack else "no"),
        ("Fallback referral fee % / FBA fee", f"{rules.referral_fee_pct} / {rules.fba_fee_default}"),
        ("Input filters (before Keepa)", prefilter.describe() if prefilter else "none"),
        ("", ""),
        ("NOTE", "est_monthly_profit = monthly sales / (new offers + 1) x profit per unit. "
                 "It assumes sales split evenly between sellers."),
        ("NOTE", "In two-pass mode, ROI and Buy Box suppression are only checked on products that passed the "
                 "other rules. Products rejected earlier show no ROI or suppression failure, even if their ROI looks low."),
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(summary, columns=["Metric", "Value"]).to_excel(writer, sheet_name="Summary", index=False)
        for name, frame in (("Passed", passed), ("Failed", failed), ("All", df)):
            frame.to_excel(writer, sheet_name=name, index=False)
        for ws in writer.book.worksheets:
            _style(ws)
            _link_asins(ws)
    return len(passed), len(df)


def save_report(rows, rules, domain_name, path, log, prefilter=None):
    try:
        return write_report(rows, rules, domain_name, path, prefilter), path
    except PermissionError:
        target = Path(path)
        alt = str(target.with_name(f"{target.stem}_{time.strftime('%H%M%S')}.xlsx"))
        log(f"Could not write {path} (is it open in Excel?). Saving to {alt} instead")
        return write_report(rows, rules, domain_name, alt, prefilter), alt


def _finish(base, result, item):
    row = {**base, **result, **item["src"]}
    profit = row.get("profit")
    if item["case_size"] and profit is not None:
        listings = row.get("listing_pack") or 1
        row["case_size"] = item["case_size"]
        row["profit_per_case"] = round(profit / listings * item["case_size"], 2)
    return row


def is_cheap(rules, buybox, two_pass):
    return bool(two_pass and buybox and (rules.check_roi or rules.check_suppression))


@dataclass
class Plan:
    df: object
    cols: dict
    basis: str
    items: list
    skipped_by_filter: dict
    qty_column: str
    asin_n: int
    code_n: int
    no_id: int


def plan(path, rules, sheet=None, overrides=None, cost_basis=None, prefilter=None, limit=None, log=print):
    df = load_catalog(path, sheet)
    log(f"Header found on row {df.attrs['header_row']} ({df.attrs['encoding']}), {len(df)} data rows")
    explicit = cost_basis in ("unit", "case")
    if not overrides:
        saved = load_profile(df)
        if saved:
            if not explicit and saved.get("cost_basis") in ("unit", "case"):
                cost_basis = saved["cost_basis"]
            overrides = saved
    if overrides:
        overrides = {f: overrides[f] for f in FIELDS if f in overrides}
    cols = detect_columns(df, overrides)
    log("Column mapping: " + ", ".join(f"{k}={v}" for k, v in cols.items()))
    if rules.check_roi and not cols["cost"]:
        raise ValueError("ROI check is on but no cost column is mapped")
    if not (cols["asin"] or cols["code"]):
        raise ValueError("Map an ASIN or UPC/EAN column first")
    basis = resolve_basis(cost_basis, cols["cost"])
    if basis == "case" and not cols["pack"]:
        if explicit:
            raise ValueError("Cost is set to per case, but no units-per-case column is mapped. "
                             "Map one, or set the cost to per unit.")
        log("The cost column looks like a per-case price but no units-per-case column is mapped, "
            "so it is treated as per unit")
        basis = "unit"
    log(f"Cost column treated as per {basis}")

    items = build_items(df, cols, basis)
    if limit:
        items = items[:limit]
    skipped_by_filter, qty_column = {}, None
    if prefilter and prefilter.active():
        patterns = [(w, _word_pattern(w)) for w in prefilter.exclude_words if _word_pattern(w)]
        if prefilter.min_qty is not None:
            qty_column = find_qty_column(df, cols, prefilter.qty_column)
            if qty_column:
                log(f"QTY column: {qty_column}")
            else:
                log("No QTY column found, so the minimum QTY filter is ignored")
        for item in items:
            reason = prefilter_reason(item, prefilter, qty_column, patterns)
            if reason:
                skipped_by_filter[item["source_row"]] = reason
        log(f"Input filters ({prefilter.describe()}) skip {len(skipped_by_filter)} of {len(items)} rows before Keepa")
    active = [i for i in items if i["source_row"] not in skipped_by_filter]
    asin_n, code_n, no_id = count_lookups(active)
    return Plan(df, cols, basis, items, skipped_by_filter, qty_column, asin_n, code_n, no_id)


def run_screening(path, rules, api_key, domain_name="GB", sheet=None, overrides=None,
                  out=None, limit=None, buybox=True, refresh=False, batch=1,  # batch is ignored: one by one
                  dry_run=False, log=print, cost_basis=None, two_pass=True,
                  stop=None, progress=None, on_row=None,
                  prefilter=None, cache_hours=24, cache_only=False,
                  on_tokens=None, name=None, planned=None, pause=None, on_wait=None):
    p = planned or plan(path, rules, sheet, overrides, cost_basis, prefilter, limit, log)
    df, cols, basis, items, skipped_by_filter = p.df, p.cols, p.basis, p.items, p.skipped_by_filter
    if not items:
        raise ValueError("The file has no product rows to screen")
    cheap = is_cheap(rules, buybox, two_pass)
    per_product = 3 if (buybox and not cheap) else 1
    log(f"{len(items)} rows: {p.asin_n} ASIN lookups, {p.code_n} barcode lookups, "
        f"{p.no_id} skipped (no usable ID), first pass about {(p.asin_n + p.code_n) * per_product} tokens")
    if dry_run:
        return None
    if not api_key and not cache_only:
        raise ValueError("No Keepa API key provided")

    domain = DOMAINS[domain_name]
    hours = 0 if refresh else cache_hours
    if cache_only:
        hours = 10 ** 6
        log("Cache-only mode: using every cached lookup regardless of age, nothing is fetched from Keepa")
    emit = on_row or (lambda row: None)
    final = {}

    bucket = TokenBucket()

    def client(with_buybox, report=progress):
        return KeepaClient(api_key, domain, buybox=with_buybox, cache_hours=hours,
                           log=log, stop=stop, progress=report, offline=cache_only, on_tokens=on_tokens,
                           pause=pause, on_wait=on_wait, bucket=bucket)

    def record(order, row):
        final[order] = row
        emit(row)

    first_rules = replace(rules, check_roi=False, check_suppression=False) if cheap else rules
    if rules.check_suppression and not buybox:
        log("Buy Box data is off, so the suppression check can only use the approximate fallback")
    lookups = {"asin": {}, "code": {}}
    for pos, item in enumerate(items):
        item["pos"] = pos
        item["base"] = {"source_row": item["source_row"], "supplier_code": item["code"],
                        "supplier_asin": item["asin"], "supplier_title": item["title"],
                        "pack_qty": item["pack"], "supplier_cost": item["cost"]}
        if item["source_row"] in skipped_by_filter:
            record((pos, 0), {**item["base"], "status": "SKIPPED",
                              "fail_reasons": "Skipped before Keepa: " + skipped_by_filter[item["source_row"]],
                              **item["src"]})
            continue
        if item["asin"]:
            lookups["asin"].setdefault(item["asin"], []).append(item)
        elif item["code"]:
            lookups["code"].setdefault(item["code"], []).append(item)
        else:
            reason = ("Invalid barcode (wrong length or check digit)" if item["bad_code"]
                      else "No valid ASIN/UPC/EAN in row")
            record((pos, 0), {**item["base"], "status": "NO_ID", "fail_reasons": reason, **item["src"]})

    early = [0]
    served = set()
    first = client(buybox and not cheap)
    second = client(True, None) if cheap else None

    def handle(kind, results):
        pending = []
        for ident, products in results.items():
            served.add((kind, ident))
            for item in lookups[kind].get(ident, []):
                if not products:
                    record((item["pos"], 0), {**item["base"], "status": "NOT_FOUND",
                                              "fail_reasons": "Not found on Keepa", **item["src"]})
                    continue
                for k, product in enumerate(products):
                    try:
                        result = evaluate(product, item["unit_cost"], first_rules, domain)
                    except Exception as exc:
                        log(f"Could not read the Keepa data for {product.get('asin') if isinstance(product, dict) else '?'} ({exc.__class__.__name__}: {exc})")
                        record((item["pos"], k), {**item["base"], "status": "NO_DATA", "asin": product.get("asin") if isinstance(product, dict) else None,
                                                  "fail_reasons": "Keepa data for this product could not be read", **item["src"]})
                        continue
                    best = None
                    if cheap and result["status"] == "PASS" and rules.check_roi:
                        best = optimistic_roi(product, item["unit_cost"], rules, domain)
                    if best is not None and best < rules.min_roi:
                        early[0] += 1
                        result = {**result, "status": "FAIL", "pass_roi": False,
                                  "fail_reasons": f"ROI at most about {best:.0f}% even at a generous price, "
                                                  f"below {rules.min_roi:g}% (rejected before the Buy Box lookup)"}
                        record((item["pos"], k), _finish(item["base"], result, item))
                    elif cheap and result["status"] == "PASS":
                        pending.append(((item["pos"], k), item, product.get("asin"), result))
                    else:
                        record((item["pos"], k), _finish(item["base"], result, item))
        if not pending:
            return
        survivors = list(dict.fromkeys(asin for _, _, asin, _ in pending))
        log(f"{', '.join(survivors)} passed the first pass, fetching Buy Box data")
        try:
            checked = second.fetch("asin", survivors)
        except (Cancelled, RuntimeError):
            for order, item, asin, first_result in pending:
                record(order, _finish(item["base"], {**first_result, "status": "STOPPED",
                                                     "fail_reasons": "Stopped before the Buy Box check"}, item))
            raise
        for order, item, asin, _ in pending:
            found = checked.get(asin) or []
            if found:
                try:
                    result = evaluate(found[0], item["unit_cost"], rules, domain)
                except Exception as exc:
                    log(f"Could not read the Keepa data for {asin} ({exc.__class__.__name__}: {exc})")
                    result = {"asin": asin, "status": "NO_DATA", "fail_reasons": "Keepa data for this product could not be read"}
            else:
                result = {"asin": asin, "status": "NO_DATA",
                          "fail_reasons": ("Buy Box data not in cache (cache-only mode)" if cache_only
                                           else "Second-pass lookup returned nothing")}
            record(order, _finish(item["base"], result, item))

    cancelled, failure = False, None
    try:
        if lookups["asin"]:
            first.fetch("asin", list(lookups["asin"]), "Lookup by ASIN",
                        lambda results: handle("asin", results))
        if lookups["code"]:
            first.fetch("code", list(lookups["code"]), "Lookup by barcode",
                        lambda results: handle("code", results))
    except Cancelled:
        cancelled = True
    except Exception as exc:
        failure = str(exc) or exc.__class__.__name__
        log(f"Stopped by an error: {failure}")
    if early[0]:
        log(f"{early[0]} products were rejected on ROI before the Buy Box lookup, saving tokens")

    if cache_only and not cancelled and not failure:
        for kind, groups in lookups.items():
            for ident, group in groups.items():
                if (kind, ident) in served:
                    continue
                for item in group:
                    record((item["pos"], 0), {**item["base"], "status": "NOT_CACHED",
                                              "fail_reasons": "Not in cache (cache-only mode, nothing fetched)",
                                              **item["src"]})

    if not final:
        if failure:
            raise RuntimeError(failure)
        raise Cancelled()
    if cancelled or failure:
        log(f"Saving the {len(final)} rows finished so far")
        finished = {order[0] for order in final}
        for item in items:
            if item["pos"] not in finished:
                final[(item["pos"], 0)] = {**item["base"], "status": "NOT_CHECKED",
                                           "fail_reasons": "Not checked yet (the run was stopped first)",
                                           **item["src"]}
    rows = [final[key] for key in sorted(final)]

    default_out = str(Path(path).with_name("screened_" + Path(path).stem + ".xlsx"))
    out = out or default_out
    if Path(out).resolve() == Path(path).resolve():
        out = default_out
        log(f"The report path was the same as the supplier file, saving to {out} instead")
    (passed, total), out = save_report(rows, rules, domain_name, out, log, prefilter)
    save_profile(df, cols, name or Path(path).stem, basis)
    log(f"Done: {passed} of {total} rows passed all enabled rules")
    return {"out": out, "passed": passed, "total": total, "cancelled": cancelled,
            "skipped": len(skipped_by_filter), "error": failure}


CACHE_COLUMNS = ["asin", "title", "brand", "ean", "upc", "monthly_sold", "bsr_now", "bsr_avg30", "bsr_avg90",
                 "new_offers", "amazon_price", "amazon_available_flag", "new_price", "buybox_price",
                 "buybox_is_amazon", "buybox_avg90", "amazon_oos_pct_90d", "rank_drops_30d", "keepa_pack_qty",
                 "fba_fee", "referral_fee_pct", "buybox_suppressed_pct_90d", "suppression_source", "hazmat",
                 "domain_id", "fetched_at"]


def _at(values, index, divide=1):
    try:
        value = values[index]
    except (TypeError, IndexError):
        return ""
    if value is None or value < 0:
        return ""
    return round(value / divide, 2) if divide != 1 else value


def _cache_row(product, fetched):
    stats = product.get("stats") or {}
    current, avg30, avg90 = stats.get("current"), stats.get("avg30"), stats.get("avg90")
    buybox = stats.get("buyBoxPrice")
    fee = (product.get("fbaFees") or {}).get("pickAndPackFee")
    suppressed, source = suppression_stats(product)
    return {
        "asin": product.get("asin"), "title": product.get("title"), "brand": product.get("brand"),
        "ean": " ".join(product.get("eanList") or []), "upc": " ".join(product.get("upcList") or []),
        "monthly_sold": product.get("monthlySold"),
        "bsr_now": _at(current, SALES_RANK), "bsr_avg30": _at(avg30, SALES_RANK), "bsr_avg90": _at(avg90, SALES_RANK),
        "new_offers": _at(current, COUNT_NEW), "amazon_price": _at(current, AMAZON, 100),
        "amazon_available_flag": product.get("availabilityAmazon"), "new_price": _at(current, NEW, 100),
        "buybox_price": "" if buybox is None or buybox < 0 else round(buybox / 100, 2),
        "buybox_is_amazon": stats.get("buyBoxIsAmazon"), "buybox_avg90": _at(avg90, BUY_BOX, 100),
        "amazon_oos_pct_90d": _at(stats.get("outOfStockPercentage90"), AMAZON),
        "rank_drops_30d": stats.get("salesRankDrops30"), "keepa_pack_qty": product.get("packageQuantity"),
        "fba_fee": "" if not fee else round(fee / 100, 2),
        "referral_fee_pct": product.get("referralFeePercentage") or product.get("referralFeePercent"),
        "buybox_suppressed_pct_90d": "" if suppressed is None else round(suppressed),
        "suppression_source": source or "", "hazmat": hazmat_info(product) or "",
        "domain_id": product.get("domainId"), "fetched_at": fetched,
    }


def export_cache(folder=None, target=None, log=print):
    folder = Path(folder) if folder else CACHE_DIR
    if not folder.exists():
        raise ValueError(f"Folder not found: {folder}")
    target = Path(target) if target else folder.parent / "keepa_cache_export.csv"
    files = sorted(folder.glob("*.json"), key=lambda p: p.stat().st_mtime)
    best, empty, unreadable = {}, 0, 0
    for path in files:
        try:
            products = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError):
            unreadable += 1
            continue
        if not products or not isinstance(products, list):
            empty += 1
            continue
        fetched = time.strftime("%Y-%m-%d %H:%M", time.localtime(path.stat().st_mtime))
        for product in products:
            if not isinstance(product, dict) or not product.get("asin"):
                continue
            row = _cache_row(product, fetched)
            old = best.get(row["asin"])
            if old and old["buybox_price"] != "" and row["buybox_price"] == "":
                continue
            best[row["asin"]] = row

    def write(path):
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.DictWriter(handle, fieldnames=CACHE_COLUMNS)
            writer.writeheader()
            writer.writerows(best.values())

    try:
        write(target)
    except PermissionError:
        target = target.with_name(f"{target.stem}_{time.strftime('%H%M%S')}.csv")
        log("The export file is open in another program, saving to a new file instead")
        write(target)
    log(f"{len(files)} cache files read: {len(best)} products, {empty} lookups with no match, {unreadable} unreadable")
    log(f"Saved to {target}")
    return target


def parse_args():
    ap = argparse.ArgumentParser(description="Screen a supplier catalog against Keepa data. "
                                             "Run with no arguments to open the web app.")
    ap.add_argument("input", nargs="?", help="Supplier catalog (.csv / .xlsx / .xls)")
    ap.add_argument("-o", "--out")
    ap.add_argument("--sheet")
    ap.add_argument("--map", action="append", default=[], metavar="FIELD=COLUMN",
                    help="Override column detection: asin|code|cost|title|pack")
    ap.add_argument("--cost-basis", choices=["auto", "unit", "case"], default="auto",
                    help="Is the cost column per unit or per case")
    ap.add_argument("--api-key", default=os.getenv("KEEPA_API_KEY"))
    ap.add_argument("--domain", default="GB", choices=list(DOMAINS))
    ap.add_argument("--limit", type=int)
    ap.add_argument("--batch-size", type=int, default=1, help=argparse.SUPPRESS)  # ignored, always one by one
    ap.add_argument("--no-buybox", action="store_true")
    ap.add_argument("--single-pass", action="store_true", help="Fetch Buy Box data for every product")
    ap.add_argument("--refresh", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--skip", nargs="*", default=[], choices=CHECKS, help="Checks to disable")
    ap.add_argument("--auto-pack", action="store_true",
                    help="Multiply the cost by the listing's pack quantity")
    ap.add_argument("--export-cache", nargs="?", const="", metavar="FOLDER",
                    help="Write the Keepa cache to keepa_cache_export.csv and exit")
    ap.add_argument("--min-qty", type=float, help="Skip rows whose QTY column is below this, before Keepa")
    ap.add_argument("--min-cost", type=float, help="Skip rows cheaper than this, before Keepa")
    ap.add_argument("--max-cost", type=float, help="Skip rows dearer than this, before Keepa")
    ap.add_argument("--exclude-words", default="",
                    help="Comma separated words; rows whose name contains one are skipped before Keepa "
                         "(end a word with * to match any word starting with it)")
    ap.add_argument("--qty-column", help="Name of the QTY column if it is not detected")
    ap.add_argument("--cache-hours", type=float, default=24, help="How old a cached lookup may be")
    ap.add_argument("--cache-only", action="store_true",
                    help="Use only cached lookups of any age, no Keepa calls and no API key needed")
    d = Rules()
    ap.add_argument("--min-sales", type=int, default=d.min_monthly_sales)
    ap.add_argument("--bsr-min", type=int, default=d.bsr_min)
    ap.add_argument("--bsr-max", type=int, default=d.bsr_max)
    ap.add_argument("--bsr-max-cv", type=float, default=d.bsr_max_cv)
    ap.add_argument("--min-offers", type=int, default=d.min_offers)
    ap.add_argument("--max-offers", type=int, default=d.max_offers)
    ap.add_argument("--min-roi", type=float, default=d.min_roi)
    ap.add_argument("--max-suppressed", type=float, default=d.max_suppressed_pct,
                    help="Reject products whose Buy Box was suppressed more than this %% of the last 90 days")
    ap.add_argument("--amazon-max-instock", type=float, default=d.amazon_max_instock_pct)
    ap.add_argument("--price-basis", choices=["conservative", "current", "avg90"], default=d.price_basis)
    ap.add_argument("--referral-pct", type=float, default=d.referral_fee_pct)
    ap.add_argument("--fba-fee", type=float, default=d.fba_fee_default)
    ap.add_argument("--prep", type=float, default=d.prep_cost)
    ap.add_argument("--inbound", type=float, default=d.inbound_ship)
    ap.add_argument("--vat", type=float, default=d.vat_rate, help="VAT as a fraction, 0.2 for 20%%")
    return ap.parse_args()


def main():
    args = parse_args()
    if args.export_cache is not None:
        try:
            export_cache(args.export_cache or None)
        except ValueError as exc:
            sys.exit(str(exc))
        return
    if not args.input:
        from app import main as serve
        serve()
        return
    rules = Rules(
        min_monthly_sales=args.min_sales, bsr_min=args.bsr_min, bsr_max=args.bsr_max,
        bsr_max_cv=args.bsr_max_cv, min_offers=args.min_offers, max_offers=args.max_offers,
        amazon_max_instock_pct=args.amazon_max_instock, min_roi=args.min_roi,
        max_suppressed_pct=args.max_suppressed,
        price_basis=args.price_basis, referral_fee_pct=args.referral_pct,
        fba_fee_default=args.fba_fee, prep_cost=args.prep, inbound_ship=args.inbound,
        vat_rate=args.vat, auto_pack=args.auto_pack, **{f"check_{c}": False for c in args.skip})
    overrides = dict(m.split("=", 1) for m in args.map if "=" in m) or None
    prefilter = PreFilter(min_qty=args.min_qty, min_cost=args.min_cost, max_cost=args.max_cost,
                          exclude_words=tuple(w.strip() for w in args.exclude_words.split(",") if w.strip()),
                          qty_column=args.qty_column)
    stopping = []

    def on_ctrl_c(*_):
        if stopping:
            raise KeyboardInterrupt
        stopping.append(1)
        print("Stopping after the current product and saving what is done (Ctrl+C again to quit now)")

    import signal
    signal.signal(signal.SIGINT, on_ctrl_c)
    try:
        result = run_screening(args.input, rules, args.api_key, args.domain, args.sheet, overrides,
                               args.out, args.limit, not args.no_buybox, args.refresh,
                               1, args.dry_run, cost_basis=args.cost_basis, stop=lambda: bool(stopping),
                               two_pass=not args.single_pass, prefilter=prefilter if prefilter.active() else None,
                               cache_hours=args.cache_hours, cache_only=args.cache_only)
    except (Cancelled, KeyboardInterrupt):
        sys.exit("Stopped")
    except (ValueError, RuntimeError) as exc:
        sys.exit(str(exc))
    except (OSError, BadZipFile) as exc:
        sys.exit(f"Could not read the file: {exc}")
    if result:
        print(f"Report saved to {result['out']}")
    if result and result.get("error"):
        sys.exit(f"Finished early: {result['error']}")


if __name__ == "__main__":
    main()
